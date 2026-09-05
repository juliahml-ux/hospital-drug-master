import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_excel():
    with open('output/115_drugs_enriched_preset.json', 'r', encoding='utf-8') as f:
        drugs = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "115年院內藥品主檔(含健保碼與價格)"

    headers = [
        "序號",
        "院內代碼",
        "健保代碼",
        "ATC碼",
        "英文商品名",
        "健保中文品名",
        "學名/成分",
        "規格劑量/單位",
        "劑型",
        "健保支付價 (NT$)",
        "健保給付規定章節",
        "給付規定PDF連結",
        "健保/食藥署查詢連結"
    ]

    # Header styling
    header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid") # Teal 600
    header_font = Font(name="微軟正黑體", size=10, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    row_font = Font(name="微軟正黑體", size=9)

    for i, d in enumerate(drugs, start=1):
        price_val = ''
        if d.get('price'):
            try:
                price_val = float(str(d['price']).replace(',', '').strip())
            except ValueError:
                price_val = str(d['price'])
        row_vals = [
            i,
            d.get('hospitalCode', ''),
            d.get('nhiCode', '') or '自費/未收載',
            d.get('atc', ''),
            d.get('brandName', ''),
            d.get('chineseName', '') or '-',
            d.get('genericName', ''),
            d.get('strength', ''),
            d.get('dosageForm', ''),
            price_val,
            d.get('ruleSection', '') or '無特殊章節',
            d.get('ruleLink', ''),
            d.get('nhiLink', '')
        ]
        ws.append(row_vals)
        current_row = i + 1
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=c)
            cell.font = row_font
            cell.border = thin_border
            if c in [1, 2, 3, 9, 10]:
                cell.alignment = align_center
            elif c == 10:
                cell.alignment = align_right
                cell.number_format = '$#,##0.00'
            else:
                cell.alignment = align_left

    # Column widths
    col_widths = [8, 14, 16, 12, 38, 30, 35, 22, 12, 18, 18, 40, 40]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output_path = 'output/115年第一季院內藥品清單_含健保碼價格給付條件.xlsx'
    wb.save(output_path)
    print(f"Successfully generated styled Excel: {output_path}")

if __name__ == '__main__':
    generate_excel()
